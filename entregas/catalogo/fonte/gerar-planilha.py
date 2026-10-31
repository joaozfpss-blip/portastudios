#!/usr/bin/env python3
"""Planilha de controle financeiro mensal da Porta Studio."""
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.data_source import NumFmt
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else '/tmp/cat/planilha.xlsx'

# ------------------------------------------------------------------ estilo
NAVY, GOLD, INK, INK2 = '1F3F8C', 'D89A2A', '17181C', '5B5E66'
BG, LINE, AUTO_FILL, INPUT_FILL = 'F4F4F1', 'DDDCD5', 'F1F1EE', 'FFF4D6'
BLUE_IN = '0000FF'
F = 'Arial'

f_title = Font(name=F, size=18, bold=True, color=INK)
f_sub = Font(name=F, size=10, color=INK2)
f_h = Font(name=F, size=10, bold=True, color='FFFFFF')
f_body = Font(name=F, size=10, color=INK)
f_in = Font(name=F, size=10, color=BLUE_IN)
f_bold = Font(name=F, size=10, bold=True, color=INK)
f_sec = Font(name=F, size=12, bold=True, color=NAVY)
fill_h = PatternFill('solid', fgColor=NAVY)
fill_in = PatternFill('solid', fgColor=INPUT_FILL)
fill_auto = PatternFill('solid', fgColor=AUTO_FILL)
fill_total = PatternFill('solid', fgColor='E4E8F3')
thin = Side(style='thin', color=LINE)
b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
b_bottom = Border(bottom=thin)
center = Alignment(horizontal='center', vertical='center')
left = Alignment(horizontal='left', vertical='center')
wrap = Alignment(horizontal='left', vertical='top', wrap_text=True)

BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"-"'
BRL0 = '"R$" #,##0;[Red]-"R$" #,##0;"-"'
USD = '#,##0.00'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE = 'dd/mm/yyyy'
MONTH = 'mmm/yy'

LANC = "'Lançamentos'"
PROJ = "'Projetos'"
RES = "'Resumo mensal'"
CF_ = "'Custos fixos'"
LST = "'Listas'"
N_LANC = 1000            # linhas de lançamento: 5..1004
LR0, LR1 = 5, 5 + N_LANC - 1
N_PROJ = 200             # projetos: 5..204
PR0, PR1 = 5, 5 + N_PROJ - 1
CAT0, CAT1 = 5, 40       # categorias em Listas


def rng(sheet, col, a=LR0, b=LR1):
    return f"{sheet}!${col}${a}:${col}${b}"


def style(c, font=f_body, fill=None, fmt=None, align=None, border=None):
    c.font = font
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = border
    return c


def inp(c, fmt=None, align=None):
    return style(c, f_in, fill_in, fmt, align, b_all)


def auto(c, fmt=None, align=None, bold=False):
    return style(c, f_bold if bold else f_body, fill_auto, fmt, align, b_all)


def header(ws, row, cols_titles, height=30):
    for col, t in cols_titles:
        c = ws[f'{col}{row}']
        c.value = t
        style(c, f_h, fill_h, align=Alignment(horizontal='center', vertical='center', wrap_text=True), border=b_all)
    ws.row_dimensions[row].height = height


def title(ws, text, sub, legend=True):
    ws['B1'] = text
    style(ws['B1'], f_title)
    ws.row_dimensions[1].height = 34
    ws['B2'] = sub
    style(ws['B2'], f_sub)
    ws.sheet_view.showGridLines = False
    if legend:
        # legenda: amarelo = você preenche · cinza = automático
        pass


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def note(cell, text):
    cell.comment = Comment(text, 'Porta Studio')
    cell.comment.width = 320
    cell.comment.height = 140


wb = Workbook()

# ===================================================================== LISTAS
ls = wb.active
ls.title = 'Listas'
title(ls, 'Listas', 'Categorias e opções usadas nos menus. Para criar uma categoria nova, escreva no fim da lista e escolha o tipo.')
widths(ls, {'A': 2, 'B': 30, 'C': 12, 'D': 9, 'E': 9, 'F': 3, 'G': 22, 'H': 3, 'I': 18, 'J': 3, 'K': 18, 'L': 3, 'M': 10, 'N': 3, 'O': 12, 'P': 3, 'Q': 8})
# Colunas: B categoria · C tipo · D nº entrada (auto) · E nº saída (auto)
header(ls, 4, [('B', 'Categoria'), ('C', 'Tipo'), ('D', 'Nº entrada'), ('E', 'Nº saída'),
               ('G', 'Forma de pagamento'), ('I', 'Tipo de projeto'), ('K', 'Status do projeto'),
               ('M', 'Moeda'), ('O', 'Cobrança'), ('Q', 'Ativo?')])
CATS = [
    ('Site WhatsApp', 'Entrada'), ('Loja virtual', 'Entrada'), ('SaaS / sistema', 'Entrada'),
    ('Mensalidade de cliente', 'Entrada'), ('Extras e ajustes', 'Entrada'),
    ('Repasse de domínio', 'Entrada'), ('Outras entradas', 'Entrada'),
    ('Inteligência artificial', 'Saída'), ('Domínios', 'Saída'), ('Hospedagem', 'Saída'),
    ('Shopify e plataformas', 'Saída'), ('Ferramentas e apps', 'Saída'),
    ('Anúncios e marketing', 'Saída'), ('Impostos (MEI / DAS)', 'Saída'),
    ('Taxas de pagamento', 'Saída'), ('Freelancers e terceiros', 'Saída'),
    ('Equipamentos', 'Saída'), ('Internet e celular', 'Saída'), ('Outras saídas', 'Saída'),
]
for i, r in enumerate(range(CAT0, CAT1 + 1)):
    cb, cc = ls[f'B{r}'], ls[f'C{r}']
    if i < len(CATS):
        cb.value, cc.value = CATS[i]
    inp(cb); inp(cc, align=center)
    ls[f'D{r}'] = f'=IF(C{r}="Entrada",COUNTIF($C${CAT0}:C{r},"Entrada"),"")'
    ls[f'E{r}'] = f'=IF(C{r}="Saída",COUNTIF($C${CAT0}:C{r},"Saída"),"")'
    auto(ls[f'D{r}'], align=center); auto(ls[f'E{r}'], align=center)
dv_tipo = DataValidation(type='list', formula1='"Entrada,Saída"', allow_blank=True)
ls.add_data_validation(dv_tipo)
dv_tipo.add(f'C{CAT0}:C{CAT1}')

SIMPLE = {
    'G': ['Pix', 'Cartão de crédito', 'Cartão de débito', 'Boleto', 'Transferência', 'Dinheiro'],
    'I': ['Site WhatsApp', 'Loja virtual', 'SaaS / sistema', 'Extras e ajustes'],
    'K': ['Orçamento', 'Em produção', 'Entregue', 'Cancelado'],
    'M': ['R$', 'US$'],
    'O': ['Mensal', 'Anual'],
    'Q': ['Sim', 'Não'],
}
for col, vals in SIMPLE.items():
    for i, v in enumerate(vals):
        inp(ls[f'{col}{5 + i}'])
        ls[f'{col}{5 + i}'] = v
    # espaço extra para formas de pagamento
    if col == 'G':
        for r in range(5 + len(vals), 15):
            inp(ls[f'G{r}'])

dn = {
    'Categorias': f"OFFSET({LST}!$B${CAT0},0,0,MAX(1,COUNTA({LST}!$B${CAT0}:$B${CAT1})),1)",
    'Formas': f"OFFSET({LST}!$G$5,0,0,MAX(1,COUNTA({LST}!$G$5:$G$14)),1)",
    'TiposProjeto': f"{LST}!$I$5:$I$8",
    'StatusProjeto': f"{LST}!$K$5:$K$8",
    'Moedas': f"{LST}!$M$5:$M$6",
    'Cobrancas': f"{LST}!$O$5:$O$6",
    'SimNao': f"{LST}!$Q$5:$Q$6",
    'ListaProjetos': f"OFFSET({PROJ}!$A${PR0},0,0,MAX(1,COUNTA({PROJ}!$A${PR0}:$A${PR1})),1)",
}
for k, v in dn.items():
    wb.defined_names[k] = DefinedName(k, attr_text=v)

# ================================================================ LANÇAMENTOS
la = wb.create_sheet('Lançamentos')
widths(la, {'A': 12, 'B': 26, 'C': 10, 'D': 46, 'E': 28, 'F': 19, 'G': 15, 'H': 10})
header(la, 4, [('A', 'Data'), ('B', 'Categoria'), ('C', 'Tipo'), ('D', 'Descrição'),
               ('E', 'Projeto (opcional)'), ('F', 'Forma de pagamento'), ('G', 'Valor (R$)'), ('H', 'Mês')])
la['A1'] = 'Lançamentos'; style(la['A1'], f_title)
la['A2'] = 'Tudo que entrou ou saiu de dinheiro, uma linha por movimento. Amarelo: você preenche. Cinza: calcula sozinho.'
style(la['A2'], f_sub)
la['A3'] = 'Escolha a categoria e o tipo (Entrada ou Saída) aparece sozinho. Digite o valor sempre positivo.'
style(la['A3'], f_sub)
for r in range(LR0, LR1 + 1):
    inp(la[f'A{r}'], DATE, center)
    inp(la[f'B{r}'])
    la[f'C{r}'] = (f'=IF(B{r}="","",IFERROR(INDEX({LST}!$C${CAT0}:$C${CAT1},'
                   f'MATCH(B{r},{LST}!$B${CAT0}:$B${CAT1},0)),"?"))')
    auto(la[f'C{r}'], align=center)
    inp(la[f'D{r}'])
    inp(la[f'E{r}'])
    inp(la[f'F{r}'])
    inp(la[f'G{r}'], BRL)
    la[f'H{r}'] = f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))'
    auto(la[f'H{r}'], MONTH, center)
la.freeze_panes = 'A5'
la.auto_filter.ref = f'A4:H{LR1}'

# exemplos
EX = [
    (dt.date(2026, 11, 3), 'Site WhatsApp', 'EXEMPLO · entrada de 50% do site (apague esta linha)',
     'EXEMPLO · Hamburgueria', 'Pix', 600),
    (dt.date(2026, 11, 5), 'Inteligência artificial', 'EXEMPLO · assinatura Claude Pro (apague esta linha)',
     None, 'Cartão de crédito', 114),
]
for i, (d, cat, desc, proj, forma, val) in enumerate(EX):
    r = LR0 + i
    la[f'A{r}'] = d; la[f'B{r}'] = cat; la[f'D{r}'] = desc
    la[f'E{r}'] = proj; la[f'F{r}'] = forma; la[f'G{r}'] = val

dv = DataValidation(type='list', formula1='=Categorias', allow_blank=True,
                    showErrorMessage=True, errorTitle='Categoria', error='Escolha uma categoria da lista (aba Listas).')
la.add_data_validation(dv); dv.add(f'B{LR0}:B{LR1}')
dv = DataValidation(type='list', formula1='=ListaProjetos', allow_blank=True, showErrorMessage=False)
la.add_data_validation(dv); dv.add(f'E{LR0}:E{LR1}')
dv = DataValidation(type='list', formula1='=Formas', allow_blank=True, showErrorMessage=False)
la.add_data_validation(dv); dv.add(f'F{LR0}:F{LR1}')
dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                    showErrorMessage=True, errorTitle='Valor',
                    error='Digite o valor sem sinal de menos. A categoria já diz se é entrada ou saída.')
la.add_data_validation(dv); dv.add(f'G{LR0}:G{LR1}')
dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                    showErrorMessage=True, errorTitle='Data', error='Digite uma data, por exemplo 05/11/2026.')
la.add_data_validation(dv); dv.add(f'A{LR0}:A{LR1}')
la.conditional_formatting.add(f'G{LR0}:G{LR1}', FormulaRule(formula=[f'$C{LR0}="Entrada"'], font=Font(name=F, color='1E7B34')))
la.conditional_formatting.add(f'G{LR0}:G{LR1}', FormulaRule(formula=[f'$C{LR0}="Saída"'], font=Font(name=F, color='B3261E')))
la.conditional_formatting.add(f'C{LR0}:C{LR1}', FormulaRule(formula=[f'$C{LR0}="?"'], fill=PatternFill('solid', fgColor='F9D7D3')))
note(la['C4'], 'Preenchido sozinho a partir da categoria (aba Listas). Se aparecer "?", a categoria digitada não existe na lista.')
note(la['E4'], 'Use o mesmo nome do projeto da aba Projetos. Assim a aba Projetos mostra quanto já recebeu e quanto falta.')
note(la['G4'], 'Sempre positivo. Ex.: 600 para uma entrada de R$ 600 ou 114 para uma assinatura de R$ 114.')

# ==================================================================== PROJETOS
pj = wb.create_sheet('Projetos')
pj['A1'] = 'Projetos e clientes'; style(pj['A1'], f_title)
pj['A2'] = 'Cada cliente fechado. O recebido e os custos vêm sozinhos dos lançamentos com o mesmo nome de projeto.'
style(pj['A2'], f_sub)
widths(pj, {'A': 28, 'B': 26, 'C': 16, 'D': 13, 'E': 12, 'F': 12, 'G': 14, 'H': 14, 'I': 14,
            'J': 14, 'K': 14, 'L': 14, 'M': 14, 'N': 34})
header(pj, 4, [('A', 'Projeto (nome único)'), ('B', 'Cliente / contato'), ('C', 'Tipo'), ('D', 'Pacote'),
               ('E', 'Fechado em'), ('F', 'Entrega prevista'), ('G', 'Status'), ('H', 'Valor fechado (R$)'),
               ('I', 'Recebido (R$)'), ('J', 'A receber (R$)'), ('K', 'Custos do projeto (R$)'),
               ('L', 'Lucro até agora (R$)'), ('M', 'Mensalidade (R$/mês)'), ('N', 'Observações')])
pj['G3'] = 'Totais'; style(pj['G3'], f_bold, align=Alignment(horizontal='right'))
for col in 'HIJKLM':
    pj[f'{col}3'] = f'=SUM({col}{PR0}:{col}{PR1})'
    auto(pj[f'{col}3'], BRL, bold=True)
    pj[f'{col}3'].fill = fill_total
for r in range(PR0, PR1 + 1):
    for col in 'ABCDN':
        inp(pj[f'{col}{r}'])
    inp(pj[f'E{r}'], DATE, center); inp(pj[f'F{r}'], DATE, center)
    inp(pj[f'G{r}'], align=center)
    inp(pj[f'H{r}'], BRL); inp(pj[f'M{r}'], BRL)
    pj[f'I{r}'] = f'=IF(A{r}="","",SUMIFS({rng(LANC, "G")},{rng(LANC, "E")},A{r},{rng(LANC, "C")},"Entrada"))'
    pj[f'J{r}'] = f'=IF(A{r}="","",IF(OR(G{r}="Em produção",G{r}="Entregue"),MAX(H{r}-I{r},0),0))'
    pj[f'K{r}'] = f'=IF(A{r}="","",SUMIFS({rng(LANC, "G")},{rng(LANC, "E")},A{r},{rng(LANC, "C")},"Saída"))'
    pj[f'L{r}'] = f'=IF(A{r}="","",I{r}-K{r})'
    for col in 'IJKL':
        auto(pj[f'{col}{r}'], BRL)
ex = ['EXEMPLO · Hamburgueria', 'Nome do dono · (62) 90000-0000', 'Site WhatsApp', 'Completo',
      dt.date(2026, 11, 3), dt.date(2026, 11, 15), 'Em produção', 1200]
for col, v in zip('ABCDEFGH', ex):
    pj[f'{col}{PR0}'] = v
pj[f'M{PR0}'] = 90
pj[f'N{PR0}'] = 'Apague esta linha quando começar'
pj.freeze_panes = 'B5'
pj.auto_filter.ref = f'A4:N{PR1}'
for col, name in (('C', 'TiposProjeto'), ('G', 'StatusProjeto')):
    dv = DataValidation(type='list', formula1=f'={name}', allow_blank=True)
    pj.add_data_validation(dv); dv.add(f'{col}{PR0}:{col}{PR1}')
pj.conditional_formatting.add(f'A{PR0}:N{PR1}', FormulaRule(formula=[f'$G{PR0}="Cancelado"'], font=Font(name=F, color='9A9A9A', strike=True)))
pj.conditional_formatting.add(f'G{PR0}:G{PR1}', FormulaRule(formula=[f'$G{PR0}="Entregue"'], fill=PatternFill('solid', fgColor='DCEFE0')))
pj.conditional_formatting.add(f'G{PR0}:G{PR1}', FormulaRule(formula=[f'$G{PR0}="Em produção"'], fill=PatternFill('solid', fgColor='E4E8F3')))
pj.conditional_formatting.add(f'J{PR0}:J{PR1}', FormulaRule(formula=[f'AND(ISNUMBER($J{PR0}),$J{PR0}>0)'], font=Font(name=F, bold=True, color='9A5B00')))
note(pj['A4'], 'Um nome por projeto, sem repetir. Use exatamente esse nome na coluna Projeto da aba Lançamentos.')
note(pj['J4'], 'Valor fechado menos o recebido. Só conta projetos Em produção ou Entregue.')
note(pj['M4'], 'Se o cliente paga plano de ajustes, manutenção ou hospedagem todo mês, coloque o valor aqui.')

# =============================================================== CUSTOS FIXOS
cf = wb.create_sheet('Custos fixos')
cf['A1'] = 'Custos fixos e assinaturas'; style(cf['A1'], f_title)
cf['A2'] = 'IA, domínio, Shopify, hospedagem: quanto o negócio custa por mês. Quando a cobrança cair, lance também em Lançamentos.'
style(cf['A2'], f_sub)
widths(cf, {'A': 34, 'B': 24, 'C': 9, 'D': 14, 'E': 11, 'F': 14, 'G': 9, 'H': 16, 'I': 16, 'J': 96})
cf['A4'] = 'Cotação do dólar (R$)'; style(cf['A4'], f_bold)
cf['B4'] = 5.50; inp(cf['B4'], '"R$" #,##0.00')
note(cf['B4'], 'Valor de referência. Atualize com a cotação do dia (busque "dólar hoje") ou com a cotação que aparece na fatura do cartão.')
cf['A5'] = 'IOF em compra no exterior'; style(cf['A5'], f_bold)
cf['B5'] = 0.035; inp(cf['B5'], '0.0%')
note(cf['B5'], 'IOF de 3,5% sobre compras internacionais no cartão (regra em vigor desde 2025). Confira na sua fatura.')
cf['G4'] = 'Total por mês'; style(cf['G4'], f_bold, align=Alignment(horizontal='right'))
cf['G5'] = 'Total por ano'; style(cf['G5'], f_bold, align=Alignment(horizontal='right'))
CF0, CF1 = 8, 37
cf['H4'] = f'=SUM(H{CF0}:H{CF1})'; auto(cf['H4'], BRL, bold=True); cf['H4'].fill = fill_total
cf['H5'] = f'=SUM(I{CF0}:I{CF1})'; auto(cf['H5'], BRL, bold=True); cf['H5'].fill = fill_total
header(cf, 7, [('A', 'Ferramenta / serviço'), ('B', 'Categoria'), ('C', 'Moeda'), ('D', 'Valor cobrado'),
               ('E', 'Cobrança'), ('F', 'Dia da cobrança'), ('G', 'Ativo?'), ('H', 'Custo por mês (R$)'),
               ('I', 'Custo por ano (R$)'), ('J', 'Observação')])
SUBS = [
    ('Claude Pro', 'Inteligência artificial', 'US$', 20, 'Mensal', None, 'Sim',
     'Plano Pro (US$ 20/mês). Confira o valor na fatura.'),
    ('ChatGPT Plus', 'Inteligência artificial', 'US$', 20, 'Mensal', None, 'Não',
     'Marque Sim se assinar.'),
    ('Gerador de imagens e vídeos com IA', 'Inteligência artificial', 'US$', None, 'Mensal', None, 'Sim',
     'Ex.: Higgsfield. Preencha com o valor do seu plano.'),
    ('Domínio .com.br (Registro.br)', 'Domínios', 'R$', 40, 'Anual', None, 'Sim',
     'Cerca de R$ 40 por ano, como diz o seu site. Uma linha por domínio que você paga.'),
    ('Netlify (hospedagem dos sites)', 'Hospedagem', 'US$', 0, 'Mensal', None, 'Sim',
     'Plano grátis. Se mudar de plano, atualize o valor.'),
    ('Shopify Basic', 'Shopify e plataformas', 'US$', 39, 'Mensal', None, 'Sim',
     'US$ 39 no plano mensal ou US$ 29/mês no anual (preço público). Se a loja for do cliente, quem paga é ele.'),
    ('Canva Pro', 'Ferramentas e apps', 'R$', None, 'Mensal', None, 'Não',
     'Preencha o valor e marque Sim se assinar.'),
    ('DAS do MEI', 'Impostos (MEI / DAS)', 'R$', None, 'Mensal', None, 'Não',
     'Se você for MEI, coloque o valor do boleto mensal.'),
    ('Internet e celular', 'Internet e celular', 'R$', None, 'Mensal', None, 'Não',
     'Só a parte que você usa para o trabalho, se quiser contar.'),
]
for i, r in enumerate(range(CF0, CF1 + 1)):
    for col in 'ABCEFGJ':
        inp(cf[f'{col}{r}'], align=center if col in 'CEFG' else None)
    inp(cf[f'D{r}'], USD)
    cf[f'H{r}'] = (f'=IF(OR(A{r}="",G{r}<>"Sim"),0,IF(C{r}="US$",D{r}*$B$4*(1+$B$5),D{r})'
                   f'/IF(E{r}="Anual",12,1))')
    cf[f'I{r}'] = f'=H{r}*12'
    auto(cf[f'H{r}'], BRL); auto(cf[f'I{r}'], BRL)
    if i < len(SUBS):
        a, b, c, d, e, f, g, j = SUBS[i]
        cf[f'A{r}'], cf[f'B{r}'], cf[f'C{r}'], cf[f'D{r}'] = a, b, c, d
        cf[f'E{r}'], cf[f'F{r}'], cf[f'G{r}'], cf[f'J{r}'] = e, f, g, j
for col, name in (('B', 'Categorias'), ('C', 'Moedas'), ('E', 'Cobrancas'), ('G', 'SimNao')):
    dv = DataValidation(type='list', formula1=f'={name}', allow_blank=True)
    cf.add_data_validation(dv); dv.add(f'{col}{CF0}:{col}{CF1}')
cf.conditional_formatting.add(f'A{CF0}:J{CF1}', FormulaRule(formula=[f'$G{CF0}="Não"'], font=Font(name=F, color='9A9A9A')))
note(cf['D7'], 'O valor como aparece na cobrança: em dólar se a moeda for US$, em reais se for R$. Valor anual se a cobrança for anual.')
note(cf['H7'], 'Convertido para reais (com IOF quando é em dólar) e dividido por 12 quando a cobrança é anual. Só conta o que está Ativo = Sim.')
cf.freeze_panes = 'A8'

# ============================================================= RESUMO MENSAL
rs = wb.create_sheet('Resumo mensal')
rs['B1'] = 'Resumo mensal'; style(rs['B1'], f_title)
rs['B2'] = 'Os 12 meses lado a lado, calculados a partir dos lançamentos. Para começar outro ano, troque o primeiro mês.'
style(rs['B2'], f_sub)
rs.sheet_view.showGridLines = False
widths(rs, {'A': 2, 'B': 26})
for c in range(3, 16):
    rs.column_dimensions[L(c)].width = 13
rs['B4'] = 'Primeiro mês'; style(rs['B4'], f_bold)
rs['C4'] = dt.date(2026, 11, 1); inp(rs['C4'], MONTH, center)
note(rs['C4'], 'Digite o primeiro dia do mês em que a planilha começa (ex.: 01/11/2026). Os 12 meses seguem a partir dele.')

M0, M1 = 7, 18
header(rs, 6, [('B', 'Mês'), ('C', 'Entradas'), ('D', 'Saídas'), ('E', 'Lucro'), ('F', 'Margem'),
               ('G', 'Lucro acumulado'), ('H', 'Nº de entradas')])
for i, r in enumerate(range(M0, M1 + 1)):
    rs[f'B{r}'] = '=DATE(YEAR($C$4),MONTH($C$4),1)' if i == 0 else f'=DATE(YEAR(B{r-1}),MONTH(B{r-1})+1,1)'
    auto(rs[f'B{r}'], MONTH, left)
    rs[f'C{r}'] = f'=SUMIFS({rng(LANC, "G")},{rng(LANC, "C")},"Entrada",{rng(LANC, "H")},$B{r})'
    rs[f'D{r}'] = f'=SUMIFS({rng(LANC, "G")},{rng(LANC, "C")},"Saída",{rng(LANC, "H")},$B{r})'
    rs[f'E{r}'] = f'=C{r}-D{r}'
    rs[f'F{r}'] = f'=IF(C{r}>0,E{r}/C{r},0)'
    rs[f'G{r}'] = f'=E{r}' if i == 0 else f'=G{r-1}+E{r}'
    rs[f'H{r}'] = f'=COUNTIFS({rng(LANC, "C")},"Entrada",{rng(LANC, "H")},$B{r})'
    for col, fmt in (('C', BRL), ('D', BRL), ('E', BRL), ('F', PCT), ('G', BRL), ('H', '0;-0;"-"')):
        auto(rs[f'{col}{r}'], fmt, center if col == 'H' else None)
T = M1 + 1
rs[f'B{T}'] = 'Total dos 12 meses'
rs[f'C{T}'] = f'=SUM(C{M0}:C{M1})'; rs[f'D{T}'] = f'=SUM(D{M0}:D{M1})'
rs[f'E{T}'] = f'=C{T}-D{T}'; rs[f'F{T}'] = f'=IF(C{T}>0,E{T}/C{T},0)'
rs[f'H{T}'] = f'=SUM(H{M0}:H{M1})'
for col, fmt in (('B', None), ('C', BRL), ('D', BRL), ('E', BRL), ('F', PCT), ('G', None), ('H', '0;-0;"-"')):
    auto(rs[f'{col}{T}'], fmt, center if col == 'H' else None, bold=True)
    rs[f'{col}{T}'].fill = fill_total


def cat_matrix(top, label, idx_col, n, tipo):
    rs[f'B{top - 1}'] = label; style(rs[f'B{top - 1}'], f_sec)
    header(rs, top, [('B', 'Categoria')] + [(L(3 + k), None) for k in range(12)] + [('O', 'Total')])
    for k in range(12):
        c = rs.cell(row=top, column=3 + k, value=f'=$B${M0 + k}')
        c.number_format = MONTH
    r0 = top + 1
    for j in range(n):
        r = r0 + j
        rs[f'B{r}'] = (f'=IFERROR(INDEX({LST}!$B${CAT0}:$B${CAT1},MATCH(ROWS($B${r0}:B{r}),'
                       f'{LST}!${idx_col}${CAT0}:${idx_col}${CAT1},0)),"")')
        auto(rs[f'B{r}'])
        for k in range(12):
            col = L(3 + k)
            rs[f'{col}{r}'] = (f'=IF($B{r}="","",SUMIFS({rng(LANC, "G")},{rng(LANC, "B")},$B{r},'
                               f'{rng(LANC, "H")},{col}${top}))')
            auto(rs[f'{col}{r}'], BRL0)
        rs[f'O{r}'] = f'=IF($B{r}="","",SUM(C{r}:N{r}))'
        auto(rs[f'O{r}'], BRL0, bold=True)
    rt = r0 + n
    rs[f'B{rt}'] = f'Total de {tipo}'
    for c in range(3, 16):
        col = L(c)
        rs[f'{col}{rt}'] = f'=SUM({col}{r0}:{col}{rt - 1})'
    for c in range(2, 16):
        auto(rs.cell(row=rt, column=c), BRL0, bold=True)
        rs.cell(row=rt, column=c).fill = fill_total
    return r0, rt


E_TOP = T + 4
e_r0, e_rt = cat_matrix(E_TOP, 'Entradas por categoria', 'D', 10, 'entradas')
S_TOP = e_rt + 4
s_r0, s_rt = cat_matrix(S_TOP, 'Saídas por categoria', 'E', 20, 'saídas')
rs.freeze_panes = 'C5'

# ====================================================================== PAINEL
pn = wb.create_sheet('Painel', 0)
pn.sheet_view.showGridLines = False
widths(pn, {'A': 2, 'B': 26, 'C': 16, 'D': 3, 'E': 26, 'F': 16, 'G': 3, 'H': 26, 'I': 16, 'J': 3, 'K': 26, 'L': 16})
pn['B1'] = 'Painel financeiro · Porta Studio'; style(pn['B1'], f_title)
pn['B2'] = 'Escolha o mês na célula amarela. Tudo aqui é calculado a partir das outras abas.'
style(pn['B2'], f_sub)
pn['B4'] = 'Mês de referência'; style(pn['B4'], f_bold, align=Alignment(vertical='center'))
pn['C4'] = dt.date(2026, 11, 1); inp(pn['C4'], MONTH, center)
pn.row_dimensions[4].height = 24
dv = DataValidation(type='list', formula1=f"={RES}!$B${M0}:$B${M1}", allow_blank=False)
pn.add_data_validation(dv); dv.add('C4')
note(pn['C4'], 'Escolha um mês da lista (são os 12 meses da aba Resumo mensal).')

MES = '$C$4'
tile_fill = PatternFill('solid', fgColor='FFFFFF')
tile_border = Border(left=thin, right=thin, top=thin, bottom=thin)


def tile(col, row, label, formula, fmt, sub, hi=False):
    c2 = L(col + 1)
    c1 = L(col)
    for rr in (row, row + 1, row + 2):
        pn.merge_cells(f'{c1}{rr}:{c2}{rr}')
    fill = PatternFill('solid', fgColor=NAVY) if hi else tile_fill
    fc, sc = ('FFFFFF', 'C9D3EC') if hi else (INK, INK2)
    a = pn[f'{c1}{row}']; a.value = label
    style(a, Font(name=F, size=10, bold=True, color=sc), fill, align=Alignment(horizontal='left', vertical='bottom', indent=1))
    b = pn[f'{c1}{row + 1}']; b.value = formula
    style(b, Font(name=F, size=20, bold=True, color=fc), fill, fmt, Alignment(horizontal='left', vertical='center', indent=1))
    s = pn[f'{c1}{row + 2}']; s.value = sub
    style(s, Font(name=F, size=9, color=sc), fill, align=Alignment(horizontal='left', vertical='top', indent=1))
    for rr in (row, row + 1, row + 2):
        for cc in (c1, c2):
            pn[f'{cc}{rr}'].fill = fill
            pn[f'{cc}{rr}'].border = Border(
                left=thin if cc == c1 else None, right=thin if cc == c2 else None,
                top=thin if rr == row else None, bottom=thin if rr == row + 2 else None)
    pn.row_dimensions[row].height = 22
    pn.row_dimensions[row + 1].height = 34
    pn.row_dimensions[row + 2].height = 20


ent = f'SUMIFS({rng(LANC, "G")},{rng(LANC, "C")},"Entrada",{rng(LANC, "H")},{MES})'
sai = f'SUMIFS({rng(LANC, "G")},{rng(LANC, "C")},"Saída",{rng(LANC, "H")},{MES})'
pn['B6'] = 'No mês escolhido'; style(pn['B6'], f_sec)
tile(2, 7, 'Entrou', f'={ent}', BRL, 'Entradas lançadas no mês')
tile(5, 7, 'Saiu', f'={sai}', BRL, 'Saídas lançadas no mês')
tile(8, 7, 'Lucro do mês', '=B8-E8', BRL, 'Entrou menos saiu', hi=True)
tile(11, 7, 'Margem', '=IF(B8>0,H8/B8,0)', PCT, 'Quanto sobra de cada R$ 100')
pn['B11'] = 'Visão geral'; style(pn['B11'], f_sec)
tile(2, 12, 'Lucro nos 12 meses', f"={RES}!E{T}", BRL, 'Período da aba Resumo mensal')
tile(5, 12, 'A receber de clientes', f"=SUM({PROJ}!J{PR0}:J{PR1})", BRL, 'Projetos em produção e entregues')
tile(8, 12, 'Custos fixos por mês', f"={CF_}!H4", BRL, 'Assinaturas ativas, em reais')
tile(11, 12, 'Mensalidades de clientes', (f'=SUMIFS({PROJ}!M{PR0}:M{PR1},{PROJ}!G{PR0}:G{PR1},"Em produção")'
                                         f'+SUMIFS({PROJ}!M{PR0}:M{PR1},{PROJ}!G{PR0}:G{PR1},"Entregue")'),
     BRL, 'Recorrente que entra todo mês')

# tabelas do mês
TB = 17
pn[f'B{TB - 1}'] = 'Para onde foi o dinheiro no mês'; style(pn[f'B{TB - 1}'], f_sec)
pn[f'E{TB - 1}'] = 'De onde veio o dinheiro no mês'; style(pn[f'E{TB - 1}'], f_sec)
header(pn, TB, [('B', 'Categoria de saída'), ('C', 'Valor'), ('E', 'Categoria de entrada'), ('F', 'Valor')], height=22)
NS, NE = 14, 8
for j in range(NS):
    r = TB + 1 + j
    pn[f'B{r}'] = (f'=IFERROR(INDEX({LST}!$B${CAT0}:$B${CAT1},MATCH(ROWS($B${TB + 1}:B{r}),'
                   f'{LST}!$E${CAT0}:$E${CAT1},0)),"")')
    pn[f'C{r}'] = f'=IF(B{r}="","",SUMIFS({rng(LANC, "G")},{rng(LANC, "B")},B{r},{rng(LANC, "H")},{MES}))'
    style(pn[f'B{r}'], f_body, border=b_bottom); style(pn[f'C{r}'], f_body, fmt=BRL, border=b_bottom)
for j in range(NE):
    r = TB + 1 + j
    pn[f'E{r}'] = (f'=IFERROR(INDEX({LST}!$B${CAT0}:$B${CAT1},MATCH(ROWS($E${TB + 1}:E{r}),'
                   f'{LST}!$D${CAT0}:$D${CAT1},0)),"")')
    pn[f'F{r}'] = f'=IF(E{r}="","",SUMIFS({rng(LANC, "G")},{rng(LANC, "B")},E{r},{rng(LANC, "H")},{MES}))'
    style(pn[f'E{r}'], f_body, border=b_bottom); style(pn[f'F{r}'], f_body, fmt=BRL, border=b_bottom)
rS = TB + 1 + NS
pn[f'B{rS}'] = 'Total'; pn[f'C{rS}'] = f'=SUM(C{TB + 1}:C{rS - 1})'
rE = TB + 1 + NE
pn[f'E{rE}'] = 'Total'; pn[f'F{rE}'] = f'=SUM(F{TB + 1}:F{rE - 1})'
for a, b_ in ((f'B{rS}', f'C{rS}'), (f'E{rE}', f'F{rE}')):
    style(pn[a], f_bold, fill_total); style(pn[b_], f_bold, fill_total, BRL)
# destaque das linhas com valor
pn.conditional_formatting.add(f'C{TB + 1}:C{rS - 1}', FormulaRule(formula=[f'AND(ISNUMBER(C{TB + 1}),C{TB + 1}>0)'], font=Font(name=F, bold=True, color='B3261E')))
pn.conditional_formatting.add(f'F{TB + 1}:F{rE - 1}', FormulaRule(formula=[f'AND(ISNUMBER(F{TB + 1}),F{TB + 1}>0)'], font=Font(name=F, bold=True, color='1E7B34')))

# gráficos
C_IN, C_OUT = '2F55B5', 'BF7F12'
ch = BarChart()
ch.type = 'col'; ch.grouping = 'clustered'
ch.title = 'Entradas e saídas por mês'
ch.y_axis.title = None; ch.x_axis.title = None
data = Reference(rs, min_col=3, max_col=4, min_row=6, max_row=M1)
cats = Reference(rs, min_col=2, min_row=M0, max_row=M1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
for s, colr in zip(ch.series, (C_IN, C_OUT)):
    s.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill=colr))
ch.gapWidth = 60; ch.overlap = -8
ch.y_axis.numFmt = NumFmt(formatCode='"R$" #,##0', sourceLinked=False)
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E6E6E1'))
ch.x_axis.numFmt = NumFmt(formatCode='mmm/yy', sourceLinked=False)
ch.y_axis.delete = False; ch.x_axis.delete = False
ch.legend.position = 'b'
ch.height = 8.2; ch.width = 17.5
pn.add_chart(ch, f'H{TB - 1}')

ch2 = BarChart()
ch2.type = 'bar'; ch2.grouping = 'clustered'
ch2.title = 'Saídas por categoria nos 12 meses'
n_out = sum(1 for _, t in CATS if t == 'Saída')
data = Reference(rs, min_col=15, min_row=s_r0, max_row=s_r0 + n_out - 1)
cats = Reference(rs, min_col=2, min_row=s_r0, max_row=s_r0 + n_out - 1)
ch2.add_data(data, titles_from_data=False)
ch2.set_categories(cats)
ch2.series[0].graphicalProperties = GraphicalProperties(solidFill=C_OUT, ln=LineProperties(solidFill=C_OUT))
ch2.gapWidth = 50
ch2.legend = None
ch2.x_axis.scaling.orientation = 'maxMin'
ch2.y_axis.numFmt = NumFmt(formatCode='"R$" #,##0', sourceLinked=False)
ch2.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E6E6E1'))
ch2.y_axis.delete = False; ch2.x_axis.delete = False
ch2.height = 8.2; ch2.width = 17.5
pn.add_chart(ch2, f'H{TB + 16}')

# ============================================================ TABELA DE PREÇOS
tp = wb.create_sheet('Tabela de preços')
tp.sheet_view.showGridLines = False
tp['B1'] = 'Tabela de preços'; style(tp['B1'], f_title)
tp['B2'] = 'Os preços do catálogo, as horas que cada pacote toma e quanto sobra por hora. Mude preço e horas para testar.'
style(tp['B2'], f_sub)
widths(tp, {'A': 2, 'B': 46, 'C': 40, 'D': 14, 'E': 12, 'F': 14, 'G': 13, 'H': 14, 'I': 14, 'J': 22})
tp['B4'] = 'Seu valor-hora mínimo (R$)'; style(tp['B4'], f_bold)
tp['C4'] = 60; inp(tp['C4'], BRL, center)
note(tp['C4'], 'O mínimo que você aceita ganhar por hora de trabalho. Pacote abaixo disso aparece em vermelho.')
tp['B5'] = 'Imposto sobre o que recebe (%)'; style(tp['B5'], f_bold)
tp['C5'] = 0; inp(tp['C5'], '0.0%', center)
note(tp['C5'], 'MEI: deixe 0% (o DAS é fixo e entra em Custos fixos). Simples Nacional ou outro regime: coloque a alíquota que o contador passar.')
header(tp, 7, [('B', 'Linha'), ('C', 'Pacote'), ('D', 'Preço (R$)'), ('E', 'Horas de trabalho'),
               ('F', 'Custos diretos (R$)'), ('G', 'Imposto (R$)'), ('H', 'Sobra (R$)'),
               ('I', 'Sobra por hora (R$)'), ('J', 'Situação')], height=34)
PRECOS = [
    ('Sites com pedido no WhatsApp', 'Essencial', 700, 8, 0),
    ('Sites com pedido no WhatsApp', 'Completo', 1200, 16, 20),
    ('Sites com pedido no WhatsApp', 'Premium', 1900, 28, 40),
    ('Lojas virtuais', 'Loja Start', 2500, 25, 0),
    ('Lojas virtuais', 'Loja Pro', 4500, 45, 0),
    ('Lojas virtuais', 'Loja Sob Medida (a partir de)', 7500, 80, 0),
    ('SaaS e sistemas', 'Sistema MVP (a partir de)', 6000, 80, 150),
    ('SaaS e sistemas', 'SaaS Completo (a partir de)', 15000, 200, 300),
    ('SaaS e sistemas', 'Plataforma + IA (a partir de)', 25000, 320, 600),
    ('Extras', 'Página extra', 250, 3, 0),
    ('Extras', 'Animação exclusiva extra', 350, 4, 0),
    ('Extras', 'Cadastro de 50 produtos', 150, 2, 0),
    ('Extras', 'Google Meu Negócio', 150, 2, 0),
    ('Extras', 'Imagens e vídeos com IA', 250, 3, 40),
    ('Extras', 'Ajuste avulso', 50, 0.5, 0),
    ('Mensalidades', 'Plano de ajustes do site (por mês)', 90, 1.5, 0),
    ('Mensalidades', 'Manutenção da loja (por mês)', 300, 5, 0),
    ('Mensalidades', 'Hospedagem e suporte de sistemas (por mês)', 400, 3, 150),
]
P0 = 8
for i, (linha, pac, preco, horas, custo) in enumerate(PRECOS):
    r = P0 + i
    tp[f'B{r}'] = linha; tp[f'C{r}'] = pac
    style(tp[f'B{r}'], f_body, border=b_all); style(tp[f'C{r}'], f_bold, border=b_all)
    tp[f'D{r}'] = preco; inp(tp[f'D{r}'], BRL)
    tp[f'E{r}'] = horas; inp(tp[f'E{r}'], '0.0', center)
    tp[f'F{r}'] = custo; inp(tp[f'F{r}'], BRL)
    tp[f'G{r}'] = f'=D{r}*$C$5'
    tp[f'H{r}'] = f'=D{r}-F{r}-G{r}'
    tp[f'I{r}'] = f'=IF(E{r}>0,H{r}/E{r},0)'
    tp[f'J{r}'] = f'=IF(I{r}>=$C$4,"Acima do mínimo","Abaixo do mínimo")'
    for col in 'GHI':
        auto(tp[f'{col}{r}'], BRL)
    auto(tp[f'J{r}'], align=center)
P1 = P0 + len(PRECOS) - 1
tp.conditional_formatting.add(f'J{P0}:J{P1}', FormulaRule(formula=[f'$J{P0}="Abaixo do mínimo"'], font=Font(name=F, bold=True, color='B3261E'), fill=PatternFill('solid', fgColor='F9D7D3')))
tp.conditional_formatting.add(f'J{P0}:J{P1}', FormulaRule(formula=[f'$J{P0}="Acima do mínimo"'], font=Font(name=F, color='1E7B34'), fill=PatternFill('solid', fgColor='DCEFE0')))
note(tp['E7'], 'Estimativa trabalhando com IA. Anote quanto tempo cada projeto levou de verdade e ajuste aqui.')
note(tp['F7'], 'O que você gasta só naquele projeto (créditos de IA, servidor durante o desenvolvimento, imagens). Assinaturas fixas ficam na aba Custos fixos.')

R = P1 + 3
tp[f'B{R}'] = 'Referências de mercado usadas para montar os preços (pesquisa de outubro de 2026)'
style(tp[f'B{R}'], f_sec)
REFS = [
    ('Landing page, freelancer intermediário', 'R$ 900 a R$ 2.500', 'bigdatauniversity.com.br/quanto-cobrar-por-uma-landing-page'),
    ('Loja Shopify ou Nuvemshop com tema e customização', 'R$ 3.000 a R$ 8.000', 'freelasemcrise.com.br/quanto-cobrar/criar-loja-virtual'),
    ('MVP de SaaS feito por freelancer', 'US$ 5 mil a US$ 20 mil', 'we0.ai (guia de MVP para startups, 2026)'),
    ('Shopify Basic (plataforma)', 'US$ 29/mês no anual · US$ 39 no mensal', 'pt.shopify.com/pricing'),
    ('Domínio .com.br', 'cerca de R$ 40 por ano', 'registro.br (valor citado no site da Porta Studio)'),
]
header(tp, R + 1, [('B', 'O quê'), ('C', 'Faixa'), ('D', 'Fonte')], height=22)
tp.merge_cells(start_row=R + 1, start_column=4, end_row=R + 1, end_column=10)
for i, (a, b_, c) in enumerate(REFS):
    r = R + 2 + i
    tp[f'B{r}'] = a; tp[f'C{r}'] = b_; tp[f'D{r}'] = c
    tp.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
    for col in 'BCD':
        style(tp[f'{col}{r}'], f_body, border=b_bottom)

# ================================================================= COMO USAR
cu = wb.create_sheet('Como usar', 0)
cu.sheet_view.showGridLines = False
widths(cu, {'A': 2, 'B': 4, 'C': 24, 'D': 90})
cu['B1'] = 'Controle financeiro · Porta Studio'; style(cu['B1'], f_title)
cu['B2'] = 'Planilha para acompanhar, mês a mês, o que entrou, o que saiu e quanto sobrou.'
style(cu['B2'], f_sub)
cu['B4'] = 'Legenda'; style(cu['B4'], f_sec)
cu['C5'] = 'Exemplo'; inp(cu['C5'])
cu['D5'] = 'Célula amarela com texto azul: você preenche.'
cu['C6'] = 'Exemplo'; auto(cu['C6'])
cu['D6'] = 'Célula cinza: calcula sozinha. Não precisa mexer.'
cu['C7'] = 'Triângulo vermelho'; style(cu['C7'], f_body)
cu['D7'] = 'Os títulos com um triângulo no canto têm uma explicação: passe o mouse ou toque neles.'
for r in (5, 6, 7):
    style(cu[f'D{r}'], f_body, align=left)
cu['B9'] = 'Passo a passo'; style(cu['B9'], f_sec)
PASSOS = [
    ('Lançamentos', 'Tudo que entrou ou saiu de dinheiro, uma linha por movimento: entrada de cliente, assinatura de IA, domínio, Shopify, anúncio. Escolha a categoria e o tipo (Entrada ou Saída) aparece sozinho. O valor é sempre positivo.'),
    ('Projetos', 'Cada cliente fechado, com o valor combinado. Use o mesmo nome do projeto nos lançamentos e a planilha mostra quanto já recebeu, quanto falta receber e quanto sobrou.'),
    ('Custos fixos', 'As assinaturas do negócio (IA, domínio, Shopify, hospedagem). Mostra quanto o negócio custa por mês, já convertendo dólar para real. Atualize a cotação de vez em quando. Quando a cobrança cair no cartão, lance também em Lançamentos.'),
    ('Painel', 'Escolha o mês na célula amarela e veja quanto entrou, saiu, o lucro, quanto falta receber e para onde foi o dinheiro.'),
    ('Resumo mensal', 'Os 12 meses lado a lado, por categoria. Para começar um ano novo, troque o primeiro mês (ou salve uma cópia do arquivo).'),
    ('Tabela de preços', 'Os preços do catálogo com as horas de cada pacote. Mostra quanto sobra por hora e avisa quando um pacote fica abaixo do seu mínimo.'),
    ('Listas', 'As categorias e as formas de pagamento dos menus. Para criar uma categoria nova, escreva no fim da lista e escolha se é Entrada ou Saída.'),
]
for i, (a, b_) in enumerate(PASSOS):
    r = 10 + i
    cu[f'B{r}'] = i + 1
    style(cu[f'B{r}'], Font(name=F, size=11, bold=True, color=NAVY), align=Alignment(horizontal='center', vertical='top'))
    cu[f'C{r}'] = a; style(cu[f'C{r}'], f_bold, align=Alignment(vertical='top'))
    cu[f'D{r}'] = b_; style(cu[f'D{r}'], f_body, align=wrap)
    cu.row_dimensions[r].height = 42
r = 10 + len(PASSOS) + 1
cu[f'B{r}'] = 'Antes de começar'; style(cu[f'B{r}'], f_sec)
DICAS = [
    'Apague as linhas marcadas com EXEMPLO nas abas Lançamentos e Projetos.',
    'Confira os valores da aba Custos fixos com as suas faturas e marque Sim só no que você assina.',
    'Lance as coisas no dia em que o dinheiro entra ou sai. No fim do mês, o Painel já está pronto.',
]
for i, t in enumerate(DICAS):
    rr = r + 1 + i
    cu[f'C{rr}'] = '•'; style(cu[f'C{rr}'], f_bold, align=Alignment(horizontal='right', vertical='top'))
    cu[f'D{rr}'] = t; style(cu[f'D{rr}'], f_body, align=wrap)

# ordem final e configurações
order = ['Como usar', 'Painel', 'Lançamentos', 'Projetos', 'Custos fixos', 'Resumo mensal', 'Tabela de preços', 'Listas']
wb._sheets = [wb[n] for n in order]
TAB = {'Como usar': GOLD, 'Painel': NAVY, 'Lançamentos': NAVY, 'Projetos': NAVY, 'Custos fixos': NAVY,
       'Resumo mensal': '5B5E66', 'Tabela de preços': '5B5E66', 'Listas': 'BBBBBB'}
for n in order:
    ws = wb[n]
    ws.sheet_properties.tabColor = TAB[n]
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT)
